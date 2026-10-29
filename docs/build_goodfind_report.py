from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

OUT = Path(r"C:\Users\Tamilarasan M\OneDrive\Desktop\GOODFIND\docs\GOODFIND_Mini_Project_Report_Revised.docx")
ASSET_DIR = OUT.parent / "goodfind_report_assets"
FOREST = "145C45"
LIGHT = "EAF2ED"
GOLD = "D9A62E"
GRAY = "5D6A63"

def font(size, bold=False):
    candidates = [
        r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()

def arrow(draw, start, end, fill=FOREST, width=4):
    fill = "#" + fill if not fill.startswith("#") else fill
    draw.line([start, end], fill=fill, width=width)
    import math
    angle = math.atan2(end[1]-start[1], end[0]-start[0])
    size = 12
    pts = [(end[0], end[1]), (end[0]-size*math.cos(angle-.5), end[1]-size*math.sin(angle-.5)), (end[0]-size*math.cos(angle+.5), end[1]-size*math.sin(angle+.5))]
    draw.polygon(pts, fill=fill)

def box(draw, xy, label, fill="EAF2ED", outline=FOREST, small=False):
    draw.rounded_rectangle(xy, radius=14, fill="#"+fill, outline="#"+outline, width=3)
    x1,y1,x2,y2=xy
    lines=label.split("\n")
    f=font(24 if not small else 20, True)
    line_h=29 if not small else 24
    total=line_h*len(lines)
    for i,line in enumerate(lines):
        bb=draw.textbbox((0,0),line,font=f)
        draw.text(((x1+x2-(bb[2]-bb[0]))/2,y1+(y2-y1-total)/2+i*line_h),line,fill="#18352C",font=f)

def make_diagrams():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    # Architecture
    im=Image.new("RGB",(1500,850),"white"); d=ImageDraw.Draw(im)
    d.text((45,30),"GOODFIND SYSTEM ARCHITECTURE",fill="#145C45",font=font(36,True))
    box(d,(70,180,390,360),"React Frontend\nPages and shared UI")
    box(d,(540,130,930,410),"FastAPI API\nAuthentication\nValidation and limits\nRoute handlers")
    box(d,(1080,90,1430,250),"PostgreSQL\nUsers, items, matches\nclaims and chat",small=True)
    box(d,(1080,330,1430,490),"Local media storage\nOpaque image paths",small=True)
    box(d,(1080,570,1430,730),"YOLO detection\nCached model\nIn-memory inference",small=True)
    arrow(d,(390,270),(540,270)); arrow(d,(930,210),(1080,170)); arrow(d,(930,285),(1080,410)); arrow(d,(930,350),(1080,650))
    d.text((555,505),"Focused services: auth  •  items  •  matching  •  claims  •  chat  •  notifications",fill="#5D6A63",font=font(22))
    im.save(ASSET_DIR/"fig_4_1_architecture.png")
    # Use case
    im=Image.new("RGB",(1500,900),"white"); d=ImageDraw.Draw(im)
    d.text((45,30),"GOODFIND USE CASES",fill="#145C45",font=font(36,True))
    box(d,(75,170,300,285),"Visitor",fill="F6F5ED")
    box(d,(75,390,300,505),"Authenticated\nuser",fill="F6F5ED")
    box(d,(75,610,300,725),"Other\nparticipant",fill="F6F5ED")
    cases=[("Browse public reports",380,130),("Register or log in",380,250),("Report lost/found item",380,370),("Review possible matches",380,490),("Submit claim or chat request",380,610),("Accept, message, block, report, close",380,730)]
    for lab,x,y in cases: box(d,(x,y,x+480,y+90),lab,small=True)
    for y in (175,435): arrow(d,(300,y),(380,175 if y<300 else 295))
    for y in (435,650,705): arrow(d,(300,y),(380,415 if y<500 else 650))
    arrow(d,(300,667),(380,775))
    im.save(ASSET_DIR/"fig_4_2_use_cases.png")
    # ER diagram
    im=Image.new("RGB",(1600,1000),"white"); d=ImageDraw.Draw(im)
    d.text((45,25),"GOODFIND DATABASE RELATIONSHIPS",fill="#145C45",font=font(36,True))
    nodes=[("users\nuser_id PK",80,120,370,220),("items\nitem_id PK\nuser_id FK",80,410,370,545),("matches\nmatch_id PK\nlost_item_id / found_item_id",520,300,880,435),("claims\nclaim_id PK\nmatch_id FK",1030,170,1380,285),("chat_requests\nrequest_id PK\nmatch_id FK",1030,390,1380,505),("conversations\nconversation_id PK\nmatch_id FK",1030,610,1380,725),("messages\nmessage_id PK\nconversation_id FK",1030,830,1380,945),("notifications\nnotification_id PK\nuser_id FK",520,650,880,785)]
    for lab,x1,y1,x2,y2 in nodes: box(d,(x1,y1,x2,y2),lab,small=True)
    arrow(d,(225,220),(225,410)); arrow(d,(370,470),(520,365)); arrow(d,(880,365),(1030,225)); arrow(d,(880,395),(1030,445)); arrow(d,(880,420),(1030,665)); arrow(d,(1205,725),(1205,830)); arrow(d,(225,220),(520,710));
    im.save(ASSET_DIR/"fig_5_1_er.png")
    # Sequence
    im=Image.new("RGB",(1600,950),"white"); d=ImageDraw.Draw(im)
    d.text((45,25),"PRIVATE CHAT REQUEST AND ACCEPTANCE",fill="#145C45",font=font(36,True))
    participants=[("User",120),("React",430),("FastAPI",760),("PostgreSQL",1100)]
    for label,x in participants:
        d.text((x,100),label,fill="#18352C",font=font(24,True)); d.line((x+40,145,x+40,850),fill="#B7C8BE",width=3)
    msgs=[("Select match",120,430), ("POST /api/chat/requests",430,760), ("Validate owner, match and ACTIVE items",760,1100), ("PENDING request + notification",1100,760), ("Accept request",120,430), ("PUT /api/chat/requests/{id}",430,760), ("Lock and revalidate lifecycle",760,1100), ("Conversation and aliases",1100,760), ("Open messages",120,430), ("GET/POST messages",430,760)]
    y=190
    for text,x1,x2 in msgs:
        d.text((x1+8,y-24),text,fill="#5D6A63",font=font(18)); arrow(d,(x1+40,y),(x2+40,y)); y+=63
    im.save(ASSET_DIR/"fig_5_2_sequence.png")
    # Activity
    im=Image.new("RGB",(1300,1200),"white"); d=ImageDraw.Draw(im)
    d.text((45,25),"REPORT, MATCH, CLAIM AND RESOLVE ACTIVITY",fill="#145C45",font=font(34,True))
    steps=["Start","Enter report details","Validate authenticated request","Save active report","Generate or refresh suggestions","Review possible match","Claim or request private chat","Accept and coordinate handover","Confirm return and close","End"]
    ys=list(range(110,1060,100))
    for i,(label,y) in enumerate(zip(steps,ys)):
        if label in ("Start","End"):
            d.ellipse((430,y,870,y+62),fill="#D9EBDD",outline="#145C45",width=3)
        else:
            d.rounded_rectangle((390,y,910,y+62),radius=14,fill="#F1F6F1",outline="#145C45",width=3)
        bb=d.textbbox((0,0),label,font=font(21,True)); d.text(((1300-(bb[2]-bb[0]))/2,y+19),label,fill="#18352C",font=font(21,True))
        if i<len(steps)-1: arrow(d,(650,y+62),(650,ys[i+1]))
    im.save(ASSET_DIR/"fig_5_3_activity.png")

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        shade(cell, FOREST)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx % 2 == 1:
                shade(cell, "F5F8F5")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text))
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 40, 35)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    return p

def add_figure(doc, image_path, caption, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(6.25))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    rr = cap.add_run(f"Fig. {number} {caption}")
    rr.bold = True
    rr.font.size = Pt(9)
    rr.font.name = "Times New Roman"
    return cap

make_diagrams()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72)
sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.82)
sec.right_margin = Inches(.82)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(30, 40, 35)
for name, size, color in (("Title", 28, "000000"), ("Heading 1", 18, "000000"), ("Heading 2", 13, "000000"), ("Heading 3", 11, "000000")):
    style = styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("GOODFIND  |  MINI PROJECT REPORT")
hr.font.size = Pt(8)
hr.font.bold = True
hr.font.color.rgb = RGBColor.from_string(GRAY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("V.S.B Engineering College  •  Fifth Semester  •  ")
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor.from_string(GRAY)
add_page_number(footer)

# Title page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(45)
r = p.add_run("GOODFIND")
r.bold = True; r.font.size = Pt(35); r.font.color.rgb = RGBColor.from_string(FOREST)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Privacy First Lost and Found Management System")
r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor.from_string(GRAY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(30)
r = p.add_run("A Mini Project Report")
r.bold = True; r.font.size = Pt(16)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Submitted in partial fulfilment of the requirements for the Fifth Semester\nAcademic Year 2026–2027").font.size = Pt(12)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(24)
r = p.add_run("Department of Computer Science and Engineering (AI&ML)")
r.bold = True; r.font.size = Pt(13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("V.S.B Engineering College, Karur").bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(30)
p.add_run("Project Guide\n").bold = True
p.add_run("Muthulakshmi\n[Guide designation to be confirmed]")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(28)
p.add_run("Project Team").bold = True
add_table(doc, ["Name", "Register Number"], [
    ("Tamilarasan M", "922524148113"), ("Sabarish S", "922524148094"),
    ("Nagulan K", "922524148068"), ("Velmurugan B", "922524148119")], [3.4, 2.1])
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(16)
p.add_run("Submission date: ____________________").italic = True
doc.add_page_break()

# Front matter templates
doc.add_heading("Bonafide Certificate", level=1)
add_body(doc, "This is to certify that the project report entitled “GOODFIND — Privacy-First Lost-and-Found Management System” is a bonafide record of the work carried out by the following students of the Department of Computer Science and Engineering (AI&ML), V.S.B Engineering College, Karur, during the academic year 2026–2027.")
add_table(doc, ["Student", "Register Number"], [("Tamilarasan M", "922524148113"), ("Sabarish S", "922524148094"), ("Nagulan K", "922524148068"), ("Velmurugan B", "922524148119")], [3.4, 2.1])
add_body(doc, "The report is submitted for academic evaluation in the Fifth Semester. Institutional wording, approval details, and signatory designations should be completed by the department before final submission.")
add_table(doc, ["Role", "Name / Designation", "Signature", "Date"], [("Project Guide", "Muthulakshmi\n[designation to be confirmed]", "________________", "__________"), ("Head of Department", "[Name and designation]", "________________", "__________"), ("Principal", "[Name and designation]", "________________", "__________")], [1.35, 2.35, 1.2, 1.1])
doc.add_page_break()

doc.add_heading("Student Declaration", level=1)
add_body(doc, "We, the undersigned students, declare that the project report entitled “GOODFIND — Privacy-First Lost-and-Found Management System” is our original academic work carried out during the Fifth Semester of the academic year 2026–2027 under the guidance of Muthulakshmi. We further declare that the report has not been submitted in whole or in part for any other academic award.")
add_body(doc, "The technical results and implementation details in this report describe the current GoodFind project as inspected at the time of preparation. Any institutional approval wording, official department naming, and registration details should be checked against college records before submission.")
add_table(doc, ["Student", "Signature", "Date"], [("Tamilarasan M\n922524148113", "________________", "__________"), ("Sabarish S\n922524148094", "________________", "__________"), ("Nagulan K\n922524148068", "________________", "__________"), ("Velmurugan B\n922524148119", "________________", "__________")], [3.4, 1.8, 1.2])
doc.add_page_break()

doc.add_heading("Acknowledgement", level=1)
add_body(doc, "We express our sincere gratitude to our project guide, Muthulakshmi, for her guidance and feedback throughout the development of GoodFind. We thank the Department of Computer Science and Engineering (AI&ML) and V.S.B Engineering College, Karur, for providing the academic environment and resources needed to complete this mini-project.")
add_body(doc, "We also thank our classmates, reviewers, and family members for their encouragement and suggestions during requirement analysis, implementation, testing, and documentation.")
doc.add_page_break()

doc.add_heading("Abstract", level=1)
add_body(doc, "GoodFind is a privacy-first lost-and-found management system designed to help people reconnect with missing belongings without requiring phone numbers, email addresses, or real names to be exchanged in public listings or private conversations. The application supports lost-item and found-item reports, public discovery, attribute-based matching, authenticated claims, anonymous chat aliases, abuse controls, notifications, item images, and optional object-detection suggestions.")
add_body(doc, "The system uses a React frontend, a Python FastAPI backend, and PostgreSQL. Authentication is based on signed bearer tokens, while authorization decisions use the authenticated user identity rather than client-supplied ownership fields. Matching produces ranked suggestions rather than proof of ownership. The current validation suite covers authentication, ownership, claims, messaging, image handling, object detection, request-size limits, notifications, and input validation.")
add_body(doc, "Keywords: lost and found, privacy, FastAPI, React, PostgreSQL, object detection, secure messaging, item matching")

doc.add_heading("Table of Contents", level=1)
for item in ["1 Introduction", "2 Existing System and Proposed System", "3 Requirements and Feasibility", "4 System Design", "5 Implementation", "6 Database Design", "7 Privacy and Security", "8 Testing and Results", "9 Limitations and Future Work", "10 Conclusion", "References"]:
    add_body(doc, item)
doc.add_page_break()
doc.add_heading("List of Figures", level=1)
for item in ["Fig. 4.1 GoodFind system architecture", "Fig. 4.2 GoodFind use cases", "Fig. 5.1 GoodFind database relationships", "Fig. 5.2 Private chat request and acceptance sequence", "Fig. 5.3 Report, match, claim and resolve activity"]:
    add_body(doc, item)
doc.add_heading("List of Tables", level=1)
for item in ["Table 2.1 Existing and proposed approach", "Table 3.1 Functional requirements", "Table 4.1 Architecture layers", "Table 6.1 Database tables", "Table 8.1 Verification results"]:
    add_body(doc, item)
doc.add_heading("List of Abbreviations", level=1)
add_table(doc, ["Abbreviation", "Meaning"], [("API", "Application Programming Interface"), ("AI", "Artificial Intelligence"), ("ASGI", "Asynchronous Server Gateway Interface"), ("CORS", "Cross-Origin Resource Sharing"), ("ER", "Entity Relationship"), ("JWT", "JSON Web Token"), ("SQL", "Structured Query Language"), ("UI", "User Interface"), ("YOLO", "You Only Look Once")], [1.55, 4.8])
doc.add_page_break()

doc.add_heading("1 Introduction", level=1)
doc.add_heading("1.1 Background", level=2)
add_body(doc, "Lost belongings are often difficult to recover because information is fragmented across informal messages, social-media posts, notice boards, and word of mouth. Existing approaches may also encourage people to publish personal contact details. GoodFind addresses this gap by keeping public reports focused on item information and providing a controlled private channel for people who have a credible possible match.")
doc.add_heading("1.2 Problem Statement", level=2)
add_body(doc, "A useful lost-and-found platform must make reports discoverable while reducing identity exposure. It must also avoid treating a similarity score as proof, prevent unauthorized access to items and conversations, validate uploaded files, and provide a clear path from a possible match to a safe handover.")
doc.add_heading("1.3 Objectives", level=2)
for text in ["Allow authenticated users to report lost and found items.", "Provide public browsing and filtering without exposing private account identifiers.", "Rank possible matches using item attributes and optional detector labels as a limited signal.", "Support claims, private one-to-one chat requests, anonymous aliases, blocking, reporting, notifications, and conversation closure.", "Protect ownership, uploads, request bodies, credentials, and message access through backend validation."]:
    add_bullet(doc, text)

doc.add_heading("2 Existing System and Proposed System", level=1)
doc.add_heading("2.1 Existing Approach", level=2)
add_body(doc, "Informal recovery workflows usually depend on public posts or direct contact. These approaches make it difficult to search consistently, verify ownership, track claims, or control who can contact whom. They can also expose phone numbers, email addresses, and personal names.")
doc.add_heading("2.2 Proposed GoodFind System", level=2)
add_body(doc, "GoodFind separates discovery from private coordination. Public item responses are restricted to report details needed for discovery. A possible match can lead to a claim or an independent private-chat request. Messaging becomes available only after request acceptance, and either participant can block, report, or close the conversation.")
add_table(doc, ["Area", "GoodFind approach"], [
    ("Identity", "Authenticated users; anonymous display aliases in chat."),
    ("Discovery", "Public lost/found listings with search, filters, pagination, and type isolation."),
    ("Matching", "Ranked suggestions from item attributes; not an ownership decision."),
    ("Coordination", "Request, acceptance, private messages, blocking, reporting, and closure."),
    ("Media", "Validated image uploads stored separately with opaque references; detection runs in memory."),
], [1.35, 5.0])

doc.add_heading("3 Requirements and Feasibility", level=1)
doc.add_heading("3.1 Functional Requirements", level=2)
for text in ["User registration and login with password hashing and bearer-token authentication.", "Lost and found report creation, updates, deletion, image attachment, and resolution.", "Public listing search by item name or description with category, location, status, and pagination filters.", "Possible-match generation and retrieval for lost and found reports.", "Claims creation and review with duplicate and terminal-state protection.", "Private chat requests, acceptance or decline, messages, notifications, blocking, abuse reports, and closure.", "Optional object detection returning labels, confidence scores, and bounding boxes."]:
    add_bullet(doc, text)
doc.add_heading("3.2 Non-Functional Requirements", level=2)
for text in ["Privacy: do not expose contact details in public listings or chat by default.", "Security: enforce authentication, ownership, participant membership, input limits, and parameterized SQL.", "Reliability: use clear validation errors, safe upload cleanup, idempotent actions, and transaction-aware lifecycle checks.", "Usability: responsive React pages with loading, empty, retry, and error states.", "Maintainability: separate backend modules for authentication, items, matching, claims, chat, notifications, detection, storage, and rate limiting."]:
    add_bullet(doc, text)
doc.add_heading("3.3 Feasibility", level=2)
add_body(doc, "The system is technically feasible with the selected stack because React supports the required forms and responsive views, FastAPI provides typed request validation and route dependencies, and PostgreSQL provides relational constraints and transaction support. The current project includes a bundled custom YOLO weight file for optional detection and a local storage path for report images.")

doc.add_heading("4 System Design", level=1)
doc.add_heading("4.1 High-Level Architecture", level=2)
add_figure(doc, ASSET_DIR / "fig_4_1_architecture.png", "GoodFind system architecture", "4.1")
add_body(doc, "The architecture separates the React presentation layer from FastAPI route handling and focused service modules. PostgreSQL is the source of truth for accounts, reports, matches, claims, conversations, messages, moderation actions, and notifications. Report photographs are stored separately using opaque references, while detection inference uses the cached model and does not persist uploaded image bytes.")
add_table(doc, ["Layer", "Components", "Responsibility"], [
    ("Presentation", "React, React Router, CSS", "Forms, listings, claims, messages, navigation, responsive states."),
    ("API", "FastAPI in backend/main.py", "Routes, authentication dependencies, validation, error mapping, CORS, body limits."),
    ("Services", "auth, items, matching, claims, chat, notifications, detection, storage", "Business rules and focused database operations."),
    ("Persistence", "PostgreSQL and item-images storage", "Users, reports, matches, claims, chat, notifications, moderation, and media references."),
], [1.2, 2.2, 2.95])
doc.add_heading("4.2 Main User Journey", level=2)
for text in ["A visitor browses public lost or found reports.", "An authenticated user submits a report with item details and may attach an image.", "The backend validates the report, stores the item, and can generate possible matches.", "The user reviews suggestions and may submit a claim or start a private chat request.", "The other participant accepts or declines a request. Accepted conversations use anonymous aliases.", "Participants coordinate a handover without needing to reveal personal contact details.", "An authorized participant can confirm return and close the conversation; closed conversations reject new messages."]:
    add_bullet(doc, text)
doc.add_heading("4.3 Use Cases", level=2)
add_figure(doc, ASSET_DIR / "fig_4_2_use_cases.png", "GoodFind use cases", "4.2")
add_body(doc, "The use-case view distinguishes public discovery from authenticated actions. A visitor can browse reports, while authenticated users can create reports, review suggestions, submit claims, and request chat. The other participant controls request acceptance and can use the moderation and closure controls available to conversation participants.")
doc.add_heading("4.4 Data Flow", level=2)
add_body(doc, "A report begins in the React form, travels through a bearer-authenticated FastAPI endpoint, and is validated before a parameterized PostgreSQL operation. Matching reads eligible records and may persist a unique match pair. Claims and chat requests reference the match, while accepted conversations reference the request and maintain participant aliases. Notifications are scoped to their recipient user.")
doc.add_heading("4.5 Matching Model", level=2)
add_body(doc, "The matching module compares relevant report attributes and produces a bounded similarity score. The current implementation uses item information such as normalized names, category, location, date proximity, description-related evidence, and a deliberately limited detection-label signal. The score ranks candidates; it does not prove that two records describe the same physical item or prove ownership. Missing information does not create positive evidence, and inactive reports are excluded from active matching workflows.")

doc.add_heading("5 Implementation", level=1)
doc.add_heading("5.1 Database and Component Design", level=2)
add_figure(doc, ASSET_DIR / "fig_5_1_er.png", "GoodFind database relationships", "5.1")
add_body(doc, "The entity relationship view shows the core identifiers and foreign-key relationships used by the current schema. Items belong to users, matches connect lost and found items, claims reference matches, and accepted chat requests create conversations with participant and message records. Notifications remain user-scoped rather than being exposed through public listings.")
doc.add_heading("5.2 Frontend", level=2)
add_body(doc, "The frontend is a Vite-powered React application. App.jsx defines public and protected routes. Shared components include SiteHeader, ProtectedRoute, ItemListing, ObjectDetectionSuggestion, and CategorySuggestions. The Home page provides reporting calls to action and links to public discovery. The listing component includes debounced, abortable requests, response-shape validation, stable pagination metadata, image fallbacks, and claim feedback. Claims and MessagesInbox provide retryable loading/error states and workflow controls.")
doc.add_heading("5.3 Backend", level=2)
add_body(doc, "FastAPI routes are collected in backend/main.py and delegate business logic to focused modules. Authentication creates signed bearer tokens and resolves the current user from the token. Item routes use the authenticated identity for ownership checks. Claims and chat services revalidate report and match lifecycle state before writes. Notifications are persisted as part of successful primary operations where applicable, and email delivery is optional and non-fatal.")
doc.add_heading("5.4 Private Chat Sequence", level=2)
add_figure(doc, ASSET_DIR / "fig_5_2_sequence.png", "Private chat request and acceptance sequence", "5.2")
add_body(doc, "The sequence shows the request and acceptance boundary. The backend validates the match, both report records, active lifecycle states, and participant roles before writing. Acceptance uses conditional updates and compatible row locking to prevent stale lifecycle changes from silently creating an ineligible conversation.")
doc.add_heading("5.5 Reporting and Resolution Activity", level=2)
add_figure(doc, ASSET_DIR / "fig_5_3_activity.png", "Report, match, claim and resolve activity", "5.3")
add_body(doc, "The activity flow is intentionally advisory: matching surfaces possibilities, and a claim or private conversation provides a controlled path for verification. Resolution is explicit and removes the completed pair from active workflows; a score or detector label does not automatically mark an item as recovered.")
doc.add_heading("5.6 Implementation Modules", level=2)
add_table(doc, ["Module", "Verified responsibility"], [
    ("auth.py", "Registration, password hashing, login, and bearer-token authentication."),
    ("items.py", "Item creation, public listings, ownership, images, and resolution."),
    ("matching.py", "Candidate retrieval, scoring, persistence, and match refresh."),
    ("claims.py", "Claim creation, duplicate handling, review, and notifications."),
    ("chat.py", "Requests, acceptance, conversations, messages, blocks, reports, and closure."),
    ("detection.py", "Validated in-memory image inference using the configured YOLO model."),
], [1.35, 5.0])
doc.add_heading("5.7 Validation and Upload Protection", level=2)
add_table(doc, ["Input", "Current protection"], [
    ("Registration", "Name 1–100, email up to 255, password 6–128 characters."),
    ("Item reports", "Names up to 200, categories up to 100, descriptions up to 5000, locations up to 255, detector labels capped at 20."),
    ("Messages", "1–2000 characters with whitespace-aware database constraint."),
    ("JSON bodies", "64 KiB application-level ASGI body limit."),
    ("Multipart bodies", "8 MiB total body limit, with the existing 5 MiB image-file limit retained."),
    ("Images", "Format validation, incremental byte limit, dimension/pixel checks, ownership checks, and safe replacement cleanup."),
], [1.45, 4.9])

doc.add_heading("6 Database Design", level=1)
add_body(doc, "The baseline schema in database/schema.sql defines users, items, matches, claims, chat requests, conversations, conversation participants, messages, user blocks, user reports, and notifications. The existing ordered migrations cover detector labels, private chat, notifications, and item image/resolution support. The configured PostgreSQL database was inspected read-only during project verification and contained these tables with the expected primary keys, foreign keys, indexes, and lifecycle checks.")
add_table(doc, ["Table", "Purpose"], [
    ("users", "Account identity and password hash."), ("items", "Lost/found reports, status, media reference, and detector labels."),
    ("matches", "Unique lost/found pairs and similarity score."), ("claims", "Claimant, found-item owner, and review status."),
    ("chat_requests", "Authenticated request and acceptance lifecycle."), ("conversations / participants", "Private membership, aliases, blocking state, and closure."),
    ("messages", "Validated private message content and timestamps."), ("notifications", "User-scoped in-app event feed."),
    ("user_blocks / user_reports", "Abuse-prevention controls and confidential reports."),
], [2.1, 4.25])
doc.add_heading("6.1 Important Constraints", level=2)
for text in ["Items restrict item_type to LOST or FOUND and lifecycle status to ACTIVE, HANDOVER_IN_PROGRESS, RESOLVED, or CLOSED.", "Chat requests restrict status to PENDING, ACCEPTED, DECLINED, or CANCELLED.", "Conversations restrict status to ACTIVE or CLOSED.", "Unique indexes prevent duplicate matches, duplicate claims per user/match, duplicate conversations, and repeated moderation reports.", "Foreign keys use cascading or restricted behavior appropriate to the relationship."]:
    add_bullet(doc, text)

doc.add_heading("7 Privacy and Security", level=1)
doc.add_heading("7.1 Authentication and Authorization", level=2)
add_body(doc, "Bearer-token authentication identifies the caller. Ownership and participant checks are performed in backend routes and service queries. Client-supplied user IDs are not trusted for ownership decisions. Public item responses omit owner identifiers and personal contact fields. Message reads, sends, closure, blocking, and reporting require conversation membership.")
doc.add_heading("7.2 Private Communication", level=2)
add_body(doc, "Chat requests are independent from claims but use the same authenticated match workflow. Once accepted, participants receive anonymous display aliases. New messages are rejected when a conversation is closed or blocked. Abuse reports validate a reason and prevent self-directed or duplicate invalid actions.")
doc.add_heading("7.3 Remaining Operational Considerations", level=2)
for text in ["Logout removes the client token; server-side token revocation is not currently persisted.", "Public item images are available only for eligible public records according to the current image route policy; deployment owners should still review whether all report photographs are appropriate for public display.", "Application-level request-size limits complement, but do not replace, upstream limits in a production reverse proxy or gateway.", "SMTP delivery is optional; missing or failed configuration must not be reported as successful delivery."]:
    add_bullet(doc, text)

doc.add_heading("8 Testing and Results", level=1)
add_body(doc, "The available backend test suite was executed using the working system Python environment. The suite completed with 98 passing tests and no failures. Python compilation completed successfully. Frontend lint and production build completed successfully.")
add_table(doc, ["Check", "Command", "Result"], [
    ("Backend pytest suite", "python -m pytest -q", "98 passed"),
    ("Backend unittest discovery", "python -m unittest discover -v", "98 tests, OK"),
    ("Python compilation", "python -m compileall -q .", "Exit code 0"),
    ("Frontend lint", "npm run lint", "Passed"),
    ("Frontend production build", "npm run build", "Passed"),
    ("PostgreSQL identity check", "Read-only current_database/current_user query", "goodfind / postgres"),
], [2.1, 2.55, 1.7])
doc.add_heading("8.1 Test Coverage Areas", level=2)
for text in ["Known and unknown-email login failures remain generic and use the dummy Argon2 verification path.", "Forged user IDs, cross-user item changes, unauthorized claims, non-participant message access, and invalid notification reads are rejected.", "Claim duplicates, terminal claims, competing reviews, chat lifecycle changes, blocking, reporting, and closed conversations are covered.", "Detection tests cover successful labels and boxes, malformed images, unsupported types, oversized uploads, dimension limits, inference failures, and unauthenticated access.", "Request-body tests cover JSON and multipart limits, multi-chunk bodies, invalid Content-Length, and 413 responses."]:
    add_bullet(doc, text)
add_body(doc, "Automated tests do not replace a manual browser walkthrough of every route at desktop and mobile sizes, real multi-user PostgreSQL concurrency testing, or live SMTP and model-service validation.")

doc.add_heading("9 Limitations and Future Work", level=1)
for text in ["Add server-side token revocation or short-lived access-token rotation if logout invalidation is required beyond client-side removal.", "Introduce a production reverse proxy or gateway with request-size, TLS, and edge rate-limit policies.", "Evaluate whether private report photographs should require authenticated access rather than public eligibility-based access.", "Add browser-level automated tests for responsive layouts, keyboard navigation, and end-to-end multi-user workflows.", "Evaluate a larger or domain-specific detector only after collecting representative data and defining an evaluation protocol; detector labels should remain advisory.", "Confirm the official department wording, guide designation, register numbers, signatories, and submission date before academic submission."]:
    add_bullet(doc, text)

doc.add_heading("10 Conclusion", level=1)
add_body(doc, "GoodFind provides a complete privacy-first workflow for reporting, discovering, matching, claiming, messaging about, and resolving lost and found items. Its architecture separates presentation, API routes, business services, and persistence, making the core features understandable and testable. The current implementation preserves anonymous communication and enforces important ownership and lifecycle boundaries while retaining practical public discovery.")
add_body(doc, "The project is suitable for academic demonstration after a final manual browser walkthrough and institutional review of the preliminary-page details. The report intentionally leaves official designations, signatures, dates, and any register-number corrections as editable placeholders rather than inventing them.")

doc.add_heading("References", level=1)
refs = [
    "GoodFind project source code: React frontend under frontend/src and FastAPI backend under backend.",
    "GoodFind database/schema.sql and database/migrations/001_detection_labels.sql through 004_item_images_and_resolution.sql.",
    "FastAPI documentation, https://fastapi.tiangolo.com/.",
    "React documentation, https://react.dev/.",
    "PostgreSQL documentation, https://www.postgresql.org/docs/.",
    "Ultralytics YOLO documentation, https://docs.ultralytics.com/.",
]
for ref in refs:
    add_bullet(doc, ref)

doc.add_page_break()
doc.add_heading("Appendices", level=1)
doc.add_heading("Appendix A Application Screenshots", level=2)
add_body(doc, "No genuine browser screenshots were available in the inspected project workspace. Before final college submission, capture the implemented GoodFind home page, public lost/found listings, report form, claims page, and private messages page from the running application and insert them here with captions. Do not substitute design mock-ups for implemented screens.")
doc.add_heading("Appendix B Selected API and Database Evidence", level=2)
add_body(doc, "The project source contains the route definitions in backend/main.py, focused service implementations in backend/auth.py, items.py, matching.py, claims.py, chat.py, notifications.py, detection.py, and storage.py, and the schema in database/schema.sql with ordered migration files. The read-only database verification used during report preparation confirmed the expected GoodFind tables, keys, foreign keys, indexes, and lifecycle constraints without modifying records.")
doc.add_heading("Appendix C Programme Outcome Mapping", level=2)
add_body(doc, "The following provisional mapping is included for departmental review. Replace or extend it to match the college-approved programme-outcome template before submission.")
add_table(doc, ["Outcome area", "GoodFind evidence"], [
    ("Engineering knowledge", "React, FastAPI, PostgreSQL, SQL constraints, file validation, and model inference are combined in one working system."),
    ("Problem analysis", "The project addresses fragmented recovery workflows and privacy risks in public lost-and-found communication."),
    ("Design and development", "Modular routes and services implement reporting, matching, claims, private chat, moderation, notifications, and closure."),
    ("Testing and verification", "Automated tests cover authentication, authorization, lifecycle changes, uploads, validation, detection, and messaging."),
    ("Ethics and privacy", "Anonymous aliases, restricted public fields, ownership checks, and abuse controls support responsible communication."),
], [1.75, 4.6])

doc.save(OUT)
print(OUT)
